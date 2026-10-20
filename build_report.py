"""Build Report.docx from REPORT.md (title page, inline figures, numbered captions).

Usage:  python build_report.py [REPORT.md] [Report.docx]

Supported Markdown subset: # / ## / ### headings, paragraphs, '- ' bullets,
'1. ' numbered items, '>' quotes, pipe tables, images, **bold**, *italic*,
`code`. A line '**Table N:** text' directly above a table becomes its caption;
an image '![Figure N: text](path)' becomes a figure with a caption below it.
Captions use real Word SEQ fields, so they renumber if you insert a new one
(select all, press F9). Figure/table numbers must be sequential in the file.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

BASE = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE / "REPORT.md"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else BASE / "Report.docx"
BODY_FONT, CODE_FONT = "Times New Roman", "Courier New"
TEXT_WIDTH_CM = 15.9
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*(?=\S)[^*\n]+?(?<=\S)\*)")
GREEK = {"sigma": "σ", "pi": "π", "tau": "τ", "phi": "φ"}


def prettify(text):
    """Typographic touches for non-code text."""
    text = text.replace("+/-", "±").replace("->", "→")
    text = text.replace(">=", "≥")
    return re.sub(r"\b(sigma|pi|tau|phi)\b", lambda m: GREEK[m.group(1)], text)


def set_run_font(run, name=BODY_FONT, size=None, bold=None, italic=None):
    """Set the font of a run (including the East-Asian slot)."""
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_text(par, text, size=None, bold=None, italic=None):
    """Add text with sub/superscripts (f_s, x^2) to a paragraph."""
    for part in re.split(r"(\b[A-Za-z]_[A-Za-z0-9]+\b|\^\d+)", prettify(text)):
        if not part:
            continue
        if re.fullmatch(r"[A-Za-z]_[A-Za-z0-9]+", part):
            r = par.add_run(part[0])
            set_run_font(r, size=size, bold=bold, italic=italic)
            r = par.add_run(part[2:])
            r.font.subscript = True
        elif part.startswith("^"):
            r = par.add_run(part[1:])
            r.font.superscript = True
        else:
            r = par.add_run(part)
        set_run_font(r, size=size, bold=bold, italic=italic)


def add_inline(par, text, size=None, bold=None, italic=None):
    """Add text with **bold**, *italic* and `code` spans."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            add_text(par, tok[2:-2], size, True, italic)
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            set_run_font(r, CODE_FONT, (size or 12) - 1.5, bold, italic)
        elif tok.startswith("*"):
            add_text(par, tok[1:-1], size, bold, True)
        else:
            add_text(par, tok, size, bold, italic)


def add_field(par, instruction, cached):
    """Insert a complex field (e.g. SEQ or PAGE) with a cached result."""
    def fld(kind):
        r = par.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
        return r
    fld("begin")
    r = par.add_run()
    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = f" {instruction} "
    r._r.append(el)
    fld("separate")
    run = par.add_run(str(cached))
    fld("end")
    return run


def shade(cell, fill):
    """Fill a table cell with a background colour."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def para_border(par, side="left", color="808080", size=18):
    """Add a border on one side of a paragraph."""
    p_pr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", "8"),
                 ("color", color)):
        b.set(qn(f"w:{k}"), v)
    borders.append(b)
    p_pr.append(borders)


def make_styles(doc):
    """Configure Normal, headings and Caption styles."""
    st = doc.styles
    normal = st["Normal"]
    normal.font.name, normal.font.size = BODY_FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before in (("Heading 1", 15, 16), ("Heading 2", 13, 12),
                               ("Heading 3", 12, 10)):
        h = st[name]
        h.font.name, h.font.size, h.font.bold = BODY_FONT, Pt(size), True
        h.font.italic = False
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        h.element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    cap = st["Caption"]
    cap.font.name, cap.font.size = BODY_FONT, Pt(10.5)
    cap.font.bold, cap.font.italic = False, False
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.space_after = Pt(8)


def add_caption(doc, label, num, text, center, keep_next):
    """Caption paragraph 'Label N: text' with a SEQ field for the number."""
    par = doc.add_paragraph(style="Caption")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.keep_with_next = keep_next
    r = par.add_run(f"{label} ")
    set_run_font(r, size=10.5, bold=True)
    r = add_field(par, f"SEQ {label} \\* ARABIC", num)
    set_run_font(r, size=10.5, bold=True)
    r = par.add_run(": ")
    set_run_font(r, size=10.5, bold=True)
    add_inline(par, text, size=10.5)
    return par


def add_table(doc, rows, big_title_page=False):
    """Add a pipe-table as a formatted Word table."""
    header, body = rows[0], rows[1:]
    n_cols = len(header)
    size = 11 if big_title_page else (10 if n_cols <= 5 else 9 if n_cols <= 8 else 8)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    weights = [max(5, min(28, max(len(r[c]) for r in rows))) for c in range(n_cols)]
    widths = [TEXT_WIDTH_CM * w / sum(weights) for w in weights]
    for i, row in enumerate(rows):
        tr = table.rows[i]._tr
        tr_pr = tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.width = Cm(widths[j])
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.keep_with_next = i < len(rows) - 1 and len(rows) < 25
            par.alignment = (WD_ALIGN_PARAGRAPH.LEFT if j == 0 or big_title_page
                             else WD_ALIGN_PARAGRAPH.CENTER)
            add_inline(par, text, size=size, bold=True if i == 0 else None)
            if i == 0:
                shade(cell, "D9D9D9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def parse_table(lines, i):
    """Collect pipe-table rows starting at line i; return (rows, next_i)."""
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_figure(doc, path, num, caption):
    """Centered picture with its caption below."""
    img = BASE / path
    with Image.open(img) as im:
        width_cm = min(15.5, im.width / 200 * 2.54)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_after = Pt(2)
    par.add_run().add_picture(str(img), width=Cm(width_cm))
    add_caption(doc, "Figure", num, caption, True, False)


def build():
    """Convert the Markdown source into the Word document."""
    lines = SRC.read_text().splitlines()
    doc = Document()
    make_styles(doc)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    props = doc.core_properties
    props.title = lines[0].lstrip("# ").strip()
    props.author = "MTE 3121 DSP Mini Project Team"
    props.comments = "Built from REPORT.md by build_report.py"

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(hp, "MTE 3121 | DSP Mini Project Report", size=9)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    set_run_font(r, size=9)
    set_run_font(add_field(fp, "PAGE", 1), size=9)

    title_page, fig_n, tab_n, i = True, 0, 0, 0
    pending_table_caption = None
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            if title_page:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                title_page = False
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            if pending_table_caption:
                tab_n += 1
                num, text = pending_table_caption
                assert num == tab_n, f"Table numbering: found {num}, expected {tab_n}"
                add_caption(doc, "Table", num, text, False, True)
                pending_table_caption = None
            add_table(doc, rows, big_title_page=title_page)
            continue
        m = re.fullmatch(r"\*\*Table (\d+):\*\* (.*)", stripped)
        if m:
            pending_table_caption = (int(m.group(1)), m.group(2))
            i += 1
            continue
        m = re.fullmatch(r"!\[Figure (\d+): (.*)\]\((.*)\)", stripped)
        if m:
            fig_n += 1
            assert int(m.group(1)) == fig_n, (
                f"Figure numbering: found {m.group(1)}, expected {fig_n}")
            add_figure(doc, m.group(3), fig_n, m.group(2))
            i += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and title_page:
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(80)
                par.paragraph_format.space_after = Pt(24)
                add_inline(par, text, size=20, bold=True)
            else:
                doc.add_heading(prettify(text), level=min(level, 3))
            i += 1
            continue
        if stripped.startswith(">"):
            par = doc.add_paragraph()
            para_border(par)  # before indent/spacing: schema element order
            par.paragraph_format.left_indent = Cm(0.6)
            par.paragraph_format.space_before = Pt(10)
            add_inline(par, stripped.lstrip("> ").strip(), size=11, italic=True)
            i += 1
            continue
        if stripped.startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            add_inline(par, stripped[2:])
            i += 1
            continue
        m = re.match(r"(\d+)\. (.*)", stripped)
        if m:
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            par.paragraph_format.space_after = Pt(3)
            add_inline(par, f"{m.group(1)}.\t{m.group(2)}")
            i += 1
            continue
        par = doc.add_paragraph()
        if title_page:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(par, stripped, size=13)
        else:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(par, stripped)
        i += 1
    doc.save(OUT)
    print(f"wrote {OUT.name}: {fig_n} figures, {tab_n} captioned tables")


if __name__ == "__main__":
    build()
